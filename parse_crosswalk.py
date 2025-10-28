#!/usr/bin/env python3
"""
Parse the BOC HCPCS Crosswalk Excel file and create JSON data files.
"""

import json
try:
    from openpyxl import load_workbook
except ImportError:
    import subprocess
    import sys
    print("Installing openpyxl...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import load_workbook

# Load the Excel file
print("Loading BOC_HCPCS_Crosswalk_2025_Q4_Complete.xlsx...")
wb = load_workbook('BOC_HCPCS_Crosswalk_2025_Q4_Complete.xlsx')

# Process Summary sheet
summary_sheet = wb['Summary']
boc_summary = {}

for row in summary_sheet.iter_rows(min_row=2, values_only=True):
    if row[0]:  # BOC Code exists
        boc_code = row[0]
        boc_summary[boc_code] = {
            'description': row[1],
            'total_codes': row[2],
            'hcpcs_codes': row[3].split(', ') if row[3] else []
        }

# Process Detailed_HCPCS_Codes sheet
detailed_sheet = wb['Detailed_HCPCS_Codes']
boc_detailed = {}

for row in detailed_sheet.iter_rows(min_row=2, values_only=True):
    if row[0]:  # BOC Code exists
        boc_code = row[0]
        if boc_code not in boc_detailed:
            boc_detailed[boc_code] = {
                'description': row[1],
                'codes': []
            }

        boc_detailed[boc_code]['codes'].append({
            'code': row[2],
            'description': row[3],
            'requires_accreditation': row[4],
            'effective_date': str(row[5]) if row[5] else None
        })

# Save JSON files
print("Saving boc_hcpcs_data.json...")
with open('boc_hcpcs_data.json', 'w') as f:
    json.dump(boc_summary, f, indent=2)

print("Saving boc_detailed_hcpcs_data.json...")
with open('boc_detailed_hcpcs_data.json', 'w') as f:
    json.dump(boc_detailed, f, indent=2)

print(f"✓ Parsed {len(boc_detailed)} BOC categories")
total_codes = sum(len(cat['codes']) for cat in boc_detailed.values())
print(f"✓ Total HCPCS codes: {total_codes}")
